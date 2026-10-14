"""Đọc và phân tích file CSV / XLSX để nhập (không lưu file công khai).

- CSV: giữ giá trị dạng chuỗi (không suy diễn kiểu) để bảo toàn mã có số 0 đầu.
- XLSX: dùng openpyxl (không cần Excel), không thực thi macro/công thức.
"""

from __future__ import annotations

import csv
import hashlib
import io

from openpyxl import load_workbook

MAX_FILE_BYTES = 20 * 1024 * 1024  # 20 MB
MAX_ROWS = 50000
MAX_COLUMNS = 200
PREVIEW_ROWS = 20

ALLOWED_EXTENSIONS = ("csv", "xlsx")


class FileParseError(ValueError):
    """Lỗi khi đọc/giải mã file tải lên."""


def checksum(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise FileParseError(
        "Không thể giải mã file. Vui lòng lưu file CSV dưới dạng UTF-8."
    )


def sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except Exception:  # noqa: BLE001
        return ","


def _normalize_rows(header: list[str], rows: list[list]) -> list[list]:
    width = len(header)
    if width > MAX_COLUMNS:
        raise FileParseError(f"Số cột vượt quá giới hạn {MAX_COLUMNS}.")
    return [row[:width] + [""] * (width - len(row)) if len(row) < width else row[:width]
            for row in rows]


def _parse_csv(raw: bytes) -> tuple[list[str], list[list], str]:
    text = _decode(raw)
    sample = text[:4096]
    delimiter = sniff_delimiter(sample)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    all_rows = list(reader)
    if not all_rows:
        raise FileParseError("File CSV trống.")
    header = [str(h).strip() for h in all_rows[0]]
    if not header or all(h == "" for h in header):
        raise FileParseError("File CSV thiếu dòng tiêu đề (header).")
    data = all_rows[1:]
    if len(data) > MAX_ROWS:
        raise FileParseError(f"Số dòng vượt quá giới hạn {MAX_ROWS}.")
    rows = _normalize_rows(header, data)
    return header, rows, delimiter


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def _parse_xlsx(raw: bytes) -> dict[str, tuple[list[str], list[list]]]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise FileParseError("File Excel không hợp lệ hoặc bị hỏng.") from exc
    result: dict[str, tuple[list[str], list[list]]] = {}
    for ws in wb.worksheets:
        iterator = ws.iter_rows(values_only=True)
        try:
            first = next(iterator)
        except StopIteration:
            first = None
        if not first:
            result[ws.title] = ([], [])
            continue
        header = [_cell_to_str(h) for h in first]
        rows: list[list] = []
        for row in iterator:
            rows.append([_cell_to_str(c) for c in row])
            if len(rows) > MAX_ROWS:
                raise FileParseError(f"Số dòng vượt quá giới hạn {MAX_ROWS}.")
        result[ws.title] = (header, _normalize_rows(header, rows))
    wb.close()
    if not result:
        raise FileParseError("File Excel không có sheet nào hợp lệ.")
    return result


def parse_preview(filename: str, raw: bytes, sheet: str | None = None) -> dict:
    if len(raw) > MAX_FILE_BYTES:
        raise FileParseError("File vượt quá giới hạn 20 MB.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise FileParseError("Chỉ hỗ trợ file .csv hoặc .xlsx.")
    cs = checksum(raw)

    if ext == "csv":
        header, rows, delimiter = _parse_csv(raw)
        sheets = None
    else:
        parsed = _parse_xlsx(raw)
        sheets = list(parsed.keys())
        if sheet:
            if sheet not in parsed:
                raise FileParseError(f"Không tìm thấy sheet '{sheet}'.")
            header, rows = parsed[sheet]
        else:
            first_sheet = sheets[0]
            header, rows = parsed[first_sheet]
        delimiter = None

    return {
        "format": ext,
        "filename": filename,
        "checksum": cs,
        "delimiter": delimiter,
        "sheets": sheets,
        "header": header,
        "columns": [{"index": i, "name": h} for i, h in enumerate(header)],
        "preview": rows[:PREVIEW_ROWS],
        "total_rows": len(rows),
    }


def read_rows(filename: str, raw: bytes, sheet: str | None = None) -> tuple[list[str], list[list], str]:
    """Trả về (header, rows, format) đầy đủ (đến MAX_ROWS) để tạo batch."""
    if len(raw) > MAX_FILE_BYTES:
        raise FileParseError("File vượt quá giới hạn 20 MB.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        header, rows, _ = _parse_csv(raw)
        return header, rows, "csv"
    parsed = _parse_xlsx(raw)
    if sheet:
        if sheet not in parsed:
            raise FileParseError(f"Không tìm thấy sheet '{sheet}'.")
        header, rows = parsed[sheet]
    else:
        header, rows = parsed[next(iter(parsed))]
    return header, rows, "xlsx"
